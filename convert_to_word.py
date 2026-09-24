import os
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

NAVY = RGBColor(27, 54, 93)       # #1B365D
STEEL = RGBColor(15, 76, 129)     # #0F4C81
DARK_GRAY = RGBColor(44, 62, 80)  # #2C3E50
BODY_COLOR = RGBColor(34, 34, 34) # #222222
MUTED = RGBColor(100, 116, 139)   # #64748B
CODE_COLOR = RGBColor(30, 41, 59) # #1E293B
HEX_HEADER_BG = "1B365D"
HEX_ROW_ALT = "F8FAFC"
HEX_CODE_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_QUOTE_BG = "F1F5F9"
HEX_QUOTE_BORDER = "0F4C81"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shd if any
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_callout_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_QUOTE_BORDER}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def set_code_box_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_HEADER_BG}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def clean_latex(text):
    # Common math symbol replacements for clean docx display
    t = text
    t = re.sub(r'\\text\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\mathbf\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\mathit\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\Delta\s*', 'Δ', t)
    t = re.sub(r'\\le\b', '≤', t)
    t = re.sub(r'\\ge\b', '≥', t)
    t = re.sub(r'\\times\b', '×', t)
    t = re.sub(r'\\cdot\b', '·', t)
    t = re.sub(r'\\pm\b', '±', t)
    t = re.sub(r'\\pi\b', 'π', t)
    t = re.sub(r'\\mu\b', 'μ', t)
    t = re.sub(r'\\to\b', '→', t)
    t = re.sub(r'\\in\b', '∈', t)
    t = re.sub(r'\\sum\b', '∑', t)
    t = re.sub(r'\\arctan\b', 'arctan', t)
    t = re.sub(r'\\max\b', 'max', t)
    t = re.sub(r'\\min\b', 'min', t)
    t = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', t)
    t = t.replace('^\\circ', '°').replace('^\\prime', "'").replace('^2', '²').replace('^3', '³')
    t = t.replace('SpO_2', 'SpO₂').replace('L^*', 'L*').replace('b^*', 'b*').replace('a^*', 'a*')
    t = t.replace('T_s', 'Ts').replace('T_d', 'Td')
    # strip residual single dollar signs
    t = re.sub(r'\$\$', '', t)
    t = re.sub(r'\$([^$]+)\$', r'\1', t)
    t = t.replace('\\', '')
    return t

def parse_inline_runs(paragraph, text, base_font_size=Pt(10.5), base_color=BODY_COLOR, is_italic=False):
    # Tokenize text into bold, italic, code, and plain
    # Pattern captures: `code`, **bold**, *italic*
    text = clean_latex(text)
    pattern = re.compile(r'(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)')
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('`') and token.endswith('`'):
            code_text = token[1:-1]
            run = paragraph.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size.pt * 0.9)
            run.font.color.rgb = STEEL
            run.font.bold = True
        elif token.startswith('**') and token.endswith('**'):
            bold_text = token[2:-2]
            run = paragraph.add_run(bold_text)
            run.font.name = 'Calibri'
            run.font.size = base_font_size
            run.font.color.rgb = base_color
            run.font.bold = True
            run.font.italic = is_italic
        elif token.startswith('*') and token.endswith('*'):
            italic_text = token[1:-1]
            run = paragraph.add_run(italic_text)
            run.font.name = 'Calibri'
            run.font.size = base_font_size
            run.font.color.rgb = base_color
            run.font.italic = True
        else:
            run = paragraph.add_run(token)
            run.font.name = 'Calibri'
            run.font.size = base_font_size
            run.font.color.rgb = base_color
            run.font.italic = is_italic

def add_header_footer(doc, title_text):
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"ITSOEH | {title_text}")
        hrun.font.name = 'Calibri'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = MUTED
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun1 = fp.add_run("Estación Diagnóstica Cardiovascular   |   ")
        frun1.font.name = 'Calibri'
        frun1.font.size = Pt(8.5)
        frun1.font.color.rgb = MUTED
        
        # Add Page Number field
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(fldSimple)
        
        frun2 = fp.add_run(" de ")
        frun2.font.name = 'Calibri'
        frun2.font.size = Pt(8.5)
        frun2.font.color.rgb = MUTED
        
        fldSimple2 = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
        fp._p.append(fldSimple2)

def build_docx(md_path, docx_path, title_header="Manual Técnico"):
    print(f"Iniciando conversión de {md_path} a {docx_path}...")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    doc = docx.Document()
    add_header_footer(doc, title_header)
    
    in_code_block = False
    code_lang = ""
    code_lines = []
    
    in_table = False
    table_lines = []
    
    in_cover = True if "MANUAL_TECNICO" in md_path else False
    cover_lines = []
    
    i = 0
    total_lines = len(lines)
    
    # Pre-parse cover if MANUAL_TECNICO
    if in_cover:
        while i < total_lines:
            line = lines[i]
            if "\\newpage" in line or (line.strip().startswith("# Contenido")):
                break
            cover_lines.append(line)
            i += 1
        
        # Render Cover Page
        # 1. School header
        p_inst = doc.add_paragraph()
        p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_inst.paragraph_format.space_before = Pt(30)
        p_inst.paragraph_format.space_after = Pt(4)
        r = p_inst.add_run("INSTITUTO TECNOLÓGICO SUPERIOR DEL OCCIDENTE DEL ESTADO DE HIDALGO")
        r.font.name = 'Calibri'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = NAVY
        
        p_div = doc.add_paragraph()
        p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_div.paragraph_format.space_before = Pt(0)
        p_div.paragraph_format.space_after = Pt(24)
        r = p_div.add_run("DIVISIÓN DE INGENIERÍA EN SISTEMAS COMPUTACIONALES")
        r.font.name = 'Calibri'
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = STEEL

        # Decorative Divider Table (Line)
        div_table = doc.add_table(rows=1, cols=1)
        div_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = div_table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, HEX_HEADER_BG)
        cell.paragraphs[0].paragraph_format.space_before = Pt(2)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(40)
        p_title.paragraph_format.space_after = Pt(10)
        r = p_title.add_run("MANUAL TÉCNICO Y DOCUMENTACIÓN DE CÓDIGO FUENTE")
        r.font.name = 'Calibri'
        r.font.size = Pt(21)
        r.font.bold = True
        r.font.color.rgb = NAVY

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(14)
        r = p_sub.add_run("Estación de Diagnóstico Cardiovascular y Calibración Fotométrica de Piel")
        r.font.name = 'Calibri'
        r.font.size = Pt(13.5)
        r.font.bold = True
        r.font.color.rgb = STEEL

        p_desc = doc.add_paragraph()
        p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_desc.paragraph_format.space_before = Pt(0)
        p_desc.paragraph_format.space_after = Pt(36)
        r = p_desc.add_run("Sistema Embebido ESP32, Procesamiento Digital de Señales (Butterworth SOS), Aplicación de Escritorio PyQt6, Web Bluetooth y Espacio Perceptual CIELab/ITA")
        r.font.name = 'Calibri'
        r.font.size = Pt(10)
        r.font.italic = True
        r.font.color.rgb = MUTED

        # Info Box
        info_table = doc.add_table(rows=1, cols=1)
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = info_table.cell(0, 0)
        c.width = Inches(5.8)
        set_cell_background(c, HEX_ROW_ALT)
        set_code_box_borders(c)
        set_cell_margins(c, top=140, bottom=140, left=200, right=200)

        cp = c.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(4)
        r = cp.add_run("PROYECTO:")
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = NAVY
        cp2 = c.add_paragraph()
        cp2.paragraph_format.space_after = Pt(8)
        r2 = cp2.add_run("Monitor Continuo No Invasivo de Signos Vitales (BPM, Presión Arterial PWA y SpO2) con Compensación de Pigmentación Cutánea")
        r2.font.size = Pt(9.5)

        cp3 = c.add_paragraph()
        cp3.paragraph_format.space_after = Pt(2)
        r3 = cp3.add_run("ASESORES:")
        r3.font.bold = True
        r3.font.size = Pt(9.5)
        r3.font.color.rgb = NAVY
        cp4 = c.add_paragraph()
        cp4.paragraph_format.space_after = Pt(8)
        r4 = cp4.add_run("• Mtra. Lorena Mendoza Guzmán\n• Mtra. Cristy Elizabeth Aguilar Ojeda")
        r4.font.size = Pt(9.5)

        cp5 = c.add_paragraph()
        cp5.paragraph_format.space_after = Pt(2)
        r5 = cp5.add_run("DESARROLLADOR / ALUMNO:")
        r5.font.bold = True
        r5.font.size = Pt(9.5)
        r5.font.color.rgb = NAVY
        cp6 = c.add_paragraph()
        cp6.paragraph_format.space_after = Pt(8)
        r6 = cp6.add_run("• Alejandro Pérez (8.° Semestre - Grupo A)")
        r6.font.size = Pt(9.5)

        # Date & break
        p_date = doc.add_paragraph()
        p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_date.paragraph_format.space_before = Pt(50)
        r = p_date.add_run("Hidalgo, México — Septiembre 2026")
        r.font.size = Pt(9.5)
        r.font.color.rgb = MUTED

        doc.add_page_break()
        if i < total_lines and "\\newpage" in lines[i]:
            i += 1

    # Main parsing loop
    while i < total_lines:
        raw_line = lines[i]
        line = raw_line.rstrip("\r\n")

        # Page breaks
        if "\\newpage" in line or '<div style="page-break-after: always;">' in line:
            doc.add_page_break()
            i += 1
            continue

        # Code block toggle
        if line.strip().startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = line.strip()[3:].strip()
                code_lines = []
            else:
                in_code_block = False
                # Render code box
                if code_lines:
                    table = doc.add_table(rows=1, cols=1)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    cell = table.cell(0, 0)
                    cell.width = Inches(6.5)
                    set_cell_background(cell, HEX_CODE_BG)
                    set_code_box_borders(cell)
                    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)

                    # Optional lang label
                    if code_lang:
                        lp = cell.paragraphs[0]
                        lp.paragraph_format.space_before = Pt(0)
                        lp.paragraph_format.space_after = Pt(2)
                        lrun = lp.add_run(f"[{code_lang.upper()}]")
                        lrun.font.name = 'Consolas'
                        lrun.font.size = Pt(8)
                        lrun.font.bold = True
                        lrun.font.color.rgb = MUTED
                        first_p = False
                    else:
                        first_p = True

                    # Code content
                    code_text = "\n".join(code_lines)
                    cp = cell.paragraphs[0] if first_p else cell.add_paragraph()
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(0)
                    cp.paragraph_format.line_spacing = Pt(11)
                    crun = cp.add_run(code_text)
                    crun.font.name = 'Consolas'
                    crun.font.size = Pt(8.2)
                    crun.font.color.rgb = CODE_COLOR

                    # spacing after code table
                    p_spacer = doc.add_paragraph()
                    p_spacer.paragraph_format.space_before = Pt(0)
                    p_spacer.paragraph_format.space_after = Pt(4)
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Table detection
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines.append(line.strip())
            i += 1
            # Check if table continues
            while i < total_lines and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            # Render Table
            rows_data = []
            for t_line in table_lines:
                # check if separator row
                if re.match(r'^\|[\s\-:|]+\|$', t_line):
                    continue
                # split cells
                parts = [c.strip() for c in t_line.strip('|').split('|')]
                rows_data.append(parts)

            if rows_data:
                num_cols = max(len(r) for r in rows_data)
                # normalize rows
                for r in rows_data:
                    while len(r) < num_cols:
                        r.append("")

                table = doc.add_table(rows=len(rows_data), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(table, color=HEX_BORDER, sz="4")

                for r_idx, row_data in enumerate(rows_data):
                    is_header = (r_idx == 0)
                    row_obj = table.rows[r_idx]
                    for c_idx, cell_value in enumerate(row_data):
                        cell = row_obj.cells[c_idx]
                        set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
                        if is_header:
                            set_cell_background(cell, HEX_HEADER_BG)
                        elif r_idx % 2 == 1:
                            set_cell_background(cell, HEX_ROW_ALT)
                        else:
                            set_cell_background(cell, "FFFFFF")

                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(0)
                        if is_header:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            clean_val = clean_latex(cell_value.replace('**', ''))
                            run = p.add_run(clean_val)
                            run.font.name = 'Calibri'
                            run.font.size = Pt(9.5)
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            parse_inline_runs(p, cell_value, base_font_size=Pt(9.0))

                # spacer after table
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(0)
                p_spacer.paragraph_format.space_after = Pt(6)

            table_lines = []
            continue

        # Horizontal Rule
        if line.strip() in ["---", "***", "___"]:
            # If in manual and is major separator
            i += 1
            continue

        # Headings
        if line.startswith("#"):
            h_match = re.match(r'^(#{1,5})\s+(.*)$', line)
            if h_match:
                level = len(h_match.group(1))
                h_text = h_match.group(2).strip()

                # In Manual Tecnico, each "# MÓDULO" starts on a fresh page
                if "MÓDULO" in h_text.upper():
                    doc.add_page_break()

                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True

                if level == 1:
                    p.paragraph_format.space_before = Pt(16)
                    p.paragraph_format.space_after = Pt(6)
                    run = p.add_run(clean_latex(h_text))
                    run.font.name = 'Calibri'
                    run.font.size = Pt(16)
                    run.font.bold = True
                    run.font.color.rgb = NAVY
                elif level == 2:
                    p.paragraph_format.space_before = Pt(12)
                    p.paragraph_format.space_after = Pt(4)
                    run = p.add_run(clean_latex(h_text))
                    run.font.name = 'Calibri'
                    run.font.size = Pt(13.5)
                    run.font.bold = True
                    run.font.color.rgb = STEEL
                elif level == 3:
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(3)
                    run = p.add_run(clean_latex(h_text))
                    run.font.name = 'Calibri'
                    run.font.size = Pt(11.5)
                    run.font.bold = True
                    run.font.color.rgb = DARK_GRAY
                else:
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(2)
                    run = p.add_run(clean_latex(h_text))
                    run.font.name = 'Calibri'
                    run.font.size = Pt(10.5)
                    run.font.bold = True
                    run.font.color.rgb = DARK_GRAY
                i += 1
                continue

        # Blockquote
        if line.startswith(">"):
            quote_text = line.lstrip("> ").strip()
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            cell.width = Inches(6.5)
            set_cell_background(cell, HEX_QUOTE_BG)
            set_callout_borders(cell)
            set_cell_margins(cell, top=70, bottom=70, left=140, right=140)
            qp = cell.paragraphs[0]
            qp.paragraph_format.space_before = Pt(0)
            qp.paragraph_format.space_after = Pt(0)
            parse_inline_runs(qp, quote_text, base_font_size=Pt(9.5), base_color=DARK_GRAY, is_italic=True)
            
            p_spacer = doc.add_paragraph()
            p_spacer.paragraph_format.space_before = Pt(0)
            p_spacer.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Lists (unordered or ordered)
        list_match = re.match(r'^(\s*)([\*\-\+]|\d+\.)\s+(.*)$', line)
        if list_match:
            indent_spaces = len(list_match.group(1))
            bullet_symbol = list_match.group(2)
            content = list_match.group(3)

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = Pt(13)

            # Left indent based on nesting
            level = indent_spaces // 2
            p.paragraph_format.left_indent = Inches(0.25 * (level + 1))

            # Marker
            if bullet_symbol in ['*', '-', '+']:
                marker = "• " if level == 0 else ("– " if level == 1 else "▪ ")
            else:
                marker = f"{bullet_symbol} "
            
            m_run = p.add_run(marker)
            m_run.font.name = 'Calibri'
            m_run.font.size = Pt(10)
            m_run.font.bold = True
            m_run.font.color.rgb = STEEL

            parse_inline_runs(p, content, base_font_size=Pt(10))
            i += 1
            continue

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = Pt(14)
        parse_inline_runs(p, line, base_font_size=Pt(10.5))
        i += 1

    # Save docx
    doc.save(docx_path)
    print(f"¡Documento Word guardado exitosamente en: {docx_path}!")

if __name__ == "__main__":
    base_dir = r"C:\Users\alepe\.gemini\antigravity\scratch\tensiometro_extract"
    artifact_dir = r"C:\Users\alepe\.gemini\antigravity\brain\a2dff3a4-2419-4df1-b0d5-ead2f08f3c8b"
    
    # 1. DOCUMENTACION.md -> DOCUMENTACION.docx
    doc_md = os.path.join(base_dir, "DOCUMENTACION.md")
    doc_docx = os.path.join(base_dir, "DOCUMENTACION.docx")
    build_docx(doc_md, doc_docx, title_header="Documentación Técnica")
    shutil.copy2(doc_docx, os.path.join(artifact_dir, "DOCUMENTACION.docx"))
    print("DOCUMENTACION.docx copiado a artefactos.")

    # 2. MANUAL_TECNICO.md -> MANUAL_TECNICO.docx
    man_md = os.path.join(base_dir, "MANUAL_TECNICO.md")
    man_docx = os.path.join(base_dir, "MANUAL_TECNICO.docx")
    build_docx(man_md, man_docx, title_header="Manual Técnico de Código")
    shutil.copy2(man_docx, os.path.join(artifact_dir, "MANUAL_TECNICO.docx"))
    print("MANUAL_TECNICO.docx copiado a artefactos.")
